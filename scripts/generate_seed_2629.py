import re
import sys
import uuid
import json
import argparse
import calendar
from datetime import datetime
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

DOC_ID = "26290000-0000-0000-0000-000000002629"
PERCENT_UNIT_ID = "d4b90eca-7f38-4600-8a26-1426dc3aa7db"
QTY_UNIT_ID = "a2765126-d188-47ec-8d70-40d2244349a3"

# Dùng đúng các cơ quan đã có sẵn trong bảng Agencies, KHÔNG thêm cơ quan mới
BCA_ID = "bd7af0f3-8179-47cb-aca1-24ca43fc480e"
ALL_AGENCIES_ID = "00000000-0000-0000-0000-000000009999"

def resolve_agency_id(lead_text):
    text = (lead_text or "").strip().lower()
    if "bộ công an" in text or "trung tâm dữ liệu quốc gia" in text:
        return BCA_ID
    return ALL_AGENCIES_ID

ISSUE_START = "2025-12-01 00:00:00"

def parse_time(time_str):
    sl = (time_str or "").strip().lower()
    if "bắt đầu triển khai từ năm" in sl:
        year = int(re.search(r"\d{4}", sl).group())
        return max(ISSUE_START, f"{year:04d}-01-01 00:00:00"), "2030-12-31 00:00:00", True
    span = re.search(r"(\d{4})\s*[-–]\s*(\d{4})", sl)
    if span:
        first, last = map(int, span.groups())
        return max(ISSUE_START, f"{first:04d}-01-01 00:00:00"), f"{last:04d}-12-31 00:00:00", False
    month = re.search(r"tháng\s+(\d+)/(\d{4})", sl)
    if month:
        month_num, year = map(int, month.groups())
        last_day = calendar.monthrange(year, month_num)[1]
        return ISSUE_START, f"{year:04d}-{month_num:02d}-{last_day:02d} 00:00:00", False
    year = re.search(r"năm\s+(\d{4})", sl)
    if year:
        year_num = int(year.group(1))
        return max(ISSUE_START, f"{year_num:04d}-01-01 00:00:00"), f"{year_num:04d}-12-31 00:00:00", False
    return ISSUE_START, "2030-12-31 00:00:00", True


def escape_sql(text):
    if text is None:
        return "NULL"
    return "'" + str(text).replace("'", "''") + "'"

parser = argparse.ArgumentParser(description="Generate QD 2629 seed from the original appendices.")
parser.add_argument('--source', type=Path, default=Path(__file__).resolve().parents[1] / 'tài liệu phát triển' / '2629_QD-TTg_682760.docx')
parser.add_argument('--output', type=Path, default=Path(__file__).resolve().with_name('seed_2629_data.sql'))
args = parser.parse_args()
doc = Document(args.source)

# Extract Goals from Table 2
t2 = doc.tables[2]
goals = []
goal_idx = 1
current_section = ""
current_group = ""

for r_idx in range(1, len(t2.rows)):
    row = t2.rows[r_idx]
    cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
    tt, title, unit, result, lead = cells[0], cells[1], cells[2], cells[3], cells[4]
    
    if tt in ['I', 'II']:
        current_section = title
        continue
    if tt in ['I.1', 'I.2', 'II.1', 'II.2']:
        current_group = title
        continue
    
    target_num = None
    m = re.search(r'(\d+(?:\.\d+)?)', result)
    if m:
        target_num = float(m.group(1))
        
    agency_id = resolve_agency_id(lead)
    
    goals.append({
        "id": str(uuid.uuid4()),
        "code": f"MT-{goal_idx:02d}",
        "title": title,
        "unit": unit,
        "target_raw": result,
        "target_num": target_num,
        "section": current_section,
        "group": current_group,
        "lead_agency_id": agency_id,
        "lead_agency_display": lead,
        "start_date": ISSUE_START if "2025-2027" in current_section else "2028-01-01 00:00:00",
        "target_year": 2027 if "2025-2027" in current_section else 2030
    })
    goal_idx += 1

print(f"Loaded {len(goals)} goals from Table 2.")

# Extract Tasks from Table 3
t3 = doc.tables[3]
tasks = []
task_idx = 1
current_task_group = ""

for r_idx in range(1, len(t3.rows)):
    row = t3.rows[r_idx]
    cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
    tt, title, lead, coord, result, time_str = cells[0], cells[1], cells[2], cells[3], cells[4], cells[5]
    
    # Merged header rows
    if (not lead and not coord and not result) or (cells[1] == cells[2] == cells[3]):
        current_task_group = title
        continue
        
    agency_id = resolve_agency_id(lead)
    start_dt, due_dt, is_ongoing = parse_time(time_str)
    
    deliv_due = due_dt.split(" ")[0] if due_dt else "2030-12-31"
    deliverables = [{
        "Id": uuid.uuid4().hex,
        "Title": result,
        "DueDate": deliv_due,
        "CurrentStatus": "NotStarted"
    }]
    
    is_general = (agency_id == ALL_AGENCIES_ID) or ("các bộ, ngành" in lead.lower())
    
    tasks.append({
        "id": str(uuid.uuid4()),
        "code": f"NV-{task_idx:02d}",
        "group": current_task_group,
        "title": title,
        "lead_agency_id": agency_id,
        "lead_agency_display": lead,
        "coord_agency_display": coord,
        "result": result,
        "time_str": time_str,
        "start_date": start_dt,
        "due_date": due_dt,
        "is_ongoing": is_ongoing,
        "is_general": is_general,
        "deliverables": deliverables
    })
    task_idx += 1

print(f"Loaded {len(tasks)} tasks from Table 3.")

# Generate SQL script
lines = []
lines.append("SET client_encoding = 'UTF8';")
lines.append("-- ===========================================================================")
lines.append("-- SEED DATA: QUYẾT ĐỊNH SỐ 2629/QĐ-TTg")
lines.append("-- Chương trình phát triển Chính phủ số")
lines.append("-- Chỉ thêm Mục tiêu (MT) và Nhiệm vụ (NV), KHÔNG thay đổi bảng Agencies")
lines.append("-- ===========================================================================")
lines.append("BEGIN;")
lines.append("")

# 1. Ensure Document 2629 exists
lines.append("-- 1. Đảm bảo Quyết định 2629 tồn tại trong bảng Documents")
lines.append(f"""INSERT INTO "Documents" (
    "Id", "DocumentNumber", "Name", "Summary", "Signer", "IssueDate", 
    "TimeResolution", "StartYear", "EndYear", "AttachmentPath", "AttachmentFilePaths", "CreatedAt"
) VALUES (
    '{DOC_ID}',
    '2629/QĐ-TTg',
    'Quyết định số 2629/QĐ-TTg ngày 01/12/2025 của Thủ tướng Chính phủ',
    'Chương trình phát triển Chính phủ số',
    'Thủ tướng Chính phủ',
    '2025-12-01 00:00:00',
    'Range',
    2025,
    2030,
    '',
    '[]'::jsonb,
    NOW()
) ON CONFLICT ("Id") DO UPDATE SET
    "DocumentNumber" = EXCLUDED."DocumentNumber",
    "Name" = EXCLUDED."Name",
    "Summary" = EXCLUDED."Summary",
    "StartYear" = EXCLUDED."StartYear",
    "EndYear" = EXCLUDED."EndYear";
""")

# 2. Clean up existing 2629 items
lines.append("-- 2. Xóa dữ liệu cũ của riêng QĐ 2629 để chạy lại idempotent an toàn")
lines.append(f"""DELETE FROM "AgencyTaskExecutions" WHERE "GoalTaskId" IN (SELECT "Id" FROM "GoalTaskItems" WHERE "DocumentId" = '{DOC_ID}');
DELETE FROM "TargetBaselines" WHERE "GoalTaskId" IN (SELECT "Id" FROM "GoalTaskItems" WHERE "DocumentId" = '{DOC_ID}');
DELETE FROM "GoalTaskItems" WHERE "DocumentId" = '{DOC_ID}';
""")

# 3. Insert Goals
lines.append(f"-- 3. INSERT {len(goals)} MỤC TIÊU (Level 1A: ItemType = 'Goal')")
for g in goals:
    g_id = g["id"]
    custom_baseline_json = "{}"
    lines.append(f"""INSERT INTO "GoalTaskItems" (
    "Id", "DocumentId", "ItemType", "Code", "Title", "Category", 
    "Section", "Group", "IsOngoing", "IsGeneralTask", "StartDate", "DueDate", 
    "LeadAgencyId", "AssignedAgencyId", "CoordinatingAgencyIds", "UnitId", "ResultTarget", 
    "EvaluationType", "CalculationMethod", "CustomBaseline", "Deliverables", 
    "AgencyDeliverables", "DynamicKPIs", "CreatedAt", "LeadAgencyDisplayText", "CoordinatingAgenciesDisplayText"
) VALUES (
    '{g_id}',
    '{DOC_ID}',
    'Goal',
    {escape_sql(g["code"])},
    {escape_sql(g["title"])},
    'Mục tiêu',
    {escape_sql(g["section"])},
    {escape_sql(g["group"])},
    FALSE,
    FALSE,
    '{g["start_date"]}',
    '{g["target_year"]}-12-31 00:00:00',
    '{g["lead_agency_id"]}',
    NULL,
    '[]'::jsonb,
    '{PERCENT_UNIT_ID}',
    {g["target_num"]},
    'Quantitative',
    'LatestValue',
    {escape_sql(custom_baseline_json)}::jsonb,
    '[]'::jsonb,
    '{{}}'::jsonb,
    '{{}}'::jsonb,
    NOW(),
    {escape_sql(g["lead_agency_display"])},
    NULL
);""")

    # Target baseline
    tb_id = str(uuid.uuid4())
    lines.append(f"""INSERT INTO "TargetBaselines" (
    "Id", "GoalTaskId", "Year", "TargetQuantity", "TargetQualitativeStatus", "IsCustomOverride", "OverrideNote"
) VALUES (
    '{tb_id}',
    '{g_id}',
    {g["target_year"]},
    {g["target_num"]},
    NULL,
    FALSE,
    {escape_sql(f"Chỉ tiêu {g['target_raw']} đến năm {g['target_year']}")}
);""")

lines.append("")
# 4. Insert Tasks
lines.append(f"-- 4. INSERT {len(tasks)} NHIỆM VỤ (Level 1B: ItemType = 'Task') KÈM SẢN PHẨM ĐẦU RA")
for t in tasks:
    t_id = t["id"]
    deliv_json = json.dumps(t["deliverables"], ensure_ascii=False)
    lines.append(f"""INSERT INTO "GoalTaskItems" (
    "Id", "DocumentId", "ItemType", "Code", "Title", "Category", 
    "Section", "Group", "IsOngoing", "IsGeneralTask", "StartDate", "DueDate", 
    "LeadAgencyId", "AssignedAgencyId", "CoordinatingAgencyIds", "UnitId", 
    "EvaluationType", "CalculationMethod", "CustomBaseline", "Deliverables", 
    "AgencyDeliverables", "DynamicKPIs", "CreatedAt", "LeadAgencyDisplayText", "CoordinatingAgenciesDisplayText"
) VALUES (
    '{t_id}',
    '{DOC_ID}',
    'Task',
    {escape_sql(t["code"])},
    {escape_sql(t["title"])},
    'Chính phủ số',
    'Danh mục nhiệm vụ',
    {escape_sql(t["group"])},
    {str(t["is_ongoing"]).upper()},
    {str(t["is_general"]).upper()},
    '{t["start_date"]}',
    '{t["due_date"]}',
    '{t["lead_agency_id"]}',
    NULL,
    '[]'::jsonb,
    NULL,
    'Qualitative',
    'LatestValue',
    '{{}}'::jsonb,
    {escape_sql(deliv_json)}::jsonb,
    '{{}}'::jsonb,
    '{{}}'::jsonb,
    NOW(),
    {escape_sql(t["lead_agency_display"])},
    {escape_sql(t["coord_agency_display"])}
);""")

lines.append("")
lines.append("COMMIT;")
lines.append("-- ======================== KẾT THÚC SEED 2629 ========================")

REFERENCE_GUARD = '''
-- Resolve reference data on the target database; IDs differ between installations.
DO $$
BEGIN
    IF NOT EXISTS (SELECT 1 FROM "Units" WHERE "Name" = '%') THEN
        RAISE EXCEPTION 'Missing percent unit (%%) in Units';
    END IF;
    IF (SELECT count(*) FROM "Agencies" WHERE "Code" = 'BCA') <> 1 THEN
        RAISE EXCEPTION 'Expected exactly one BCA agency';
    END IF;
END $$;
'''

sql_content = "\n".join(lines)
# Reference IDs are local to each installation. Resolve by dictionary meaning/code.
sql_content = sql_content.replace(
    f"'{PERCENT_UNIT_ID}'",
    '(SELECT "Id" FROM "Units" WHERE "Name" = \'%\' ORDER BY "Id" LIMIT 1)',
).replace(
    f"'{BCA_ID}'", '(SELECT "Id" FROM "Agencies" WHERE "Code" = \'BCA\')',
)
sql_content = sql_content.replace("BEGIN;", "BEGIN;\n\n" + REFERENCE_GUARD)

sql_out_path = args.output
with open(sql_out_path, 'w', encoding='utf-8') as f:
    f.write(sql_content)

print(f"SQL file written successfully to: {sql_out_path} ({len(sql_content)} bytes)")
